from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_THEME_COLOR

OUT = r"C:\Users\Acer\Desktop\Synora-Auction-Pitch-Deck.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(5, 10, 25)
PANEL = RGBColor(15, 24, 48)
PANEL2 = RGBColor(22, 34, 62)
CYAN = RGBColor(39, 199, 226)
MINT = RGBColor(67, 226, 164)
WHITE = RGBColor(245, 248, 255)
MUTED = RGBColor(167, 181, 207)
RED = RGBColor(255, 111, 133)
GOLD = RGBColor(247, 196, 78)


def fill_bg(slide):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    # atmospheric accent bars
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    shape.fill.solid(); shape.fill.fore_color.rgb = CYAN; shape.line.fill.background()
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.42), Inches(13.333), Inches(0.08))
    shape.fill.solid(); shape.fill.fore_color.rgb = MINT; shape.line.fill.background()


def text(slide, x, y, w, h, value, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = value
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    return box


def title(slide, kicker, heading, sub=None):
    text(slide, 0.65, 0.38, 12, 0.25, kicker.upper(), 10, CYAN, True)
    text(slide, 0.65, 0.7, 12, 0.65, heading, 30, WHITE, True)
    if sub:
        text(slide, 0.68, 1.42, 11.5, 0.42, sub, 13, MUTED)


def panel(slide, x, y, w, h, color=PANEL, radius=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = color
    shape.line.color.rgb = RGBColor(38, 57, 91); shape.line.width = Pt(0.7)
    return shape


def bullet_list(slide, x, y, w, items, size=16, color=WHITE, gap=0.44):
    for i, item in enumerate(items):
        text(slide, x, y + i * gap, 0.22, 0.25, "•", size, CYAN, True)
        text(slide, x + 0.28, y + i * gap, w - 0.28, 0.34, item, size, color)


def metric(slide, x, y, value, label, accent=CYAN):
    panel(slide, x, y, 2.65, 1.15, PANEL2)
    text(slide, x + 0.2, y + 0.16, 2.2, 0.42, value, 25, accent, True)
    text(slide, x + 0.2, y + 0.68, 2.2, 0.24, label.upper(), 9, MUTED, True)


def add_footer(slide, number):
    text(slide, 12.2, 7.1, 0.55, 0.2, f"{number:02d}", 10, MUTED, True, PP_ALIGN.RIGHT)
    text(slide, 0.65, 7.1, 3.2, 0.2, "SYNORA  /  REAL-TIME AUCTIONS", 9, MUTED, True)

# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6]); fill_bg(slide)
text(slide, 0.75, 0.75, 3, 0.3, "SYNORA", 13, CYAN, True)
text(slide, 0.75, 1.45, 11.5, 1.35, "Auctions that stay\ncorrect under pressure.", 40, WHITE, True)
text(slide, 0.8, 3.2, 9.8, 0.7, "A real-time bidding system built around one hard promise: the server decides the winner, every time.", 20, MUTED)
panel(slide, 0.8, 4.65, 11.7, 1.25, PANEL)
text(slide, 1.1, 4.92, 2.5, 0.3, "THE PITCH", 10, CYAN, True)
text(slide, 1.1, 5.28, 10.7, 0.35, "Live updates for people. Serialized truth for the auction.", 22, WHITE, True)
text(slide, 0.8, 6.55, 8, 0.3, "Hackathon presentation  /  Syntax Error", 12, MUTED)
add_footer(slide, 1)

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6]); fill_bg(slide)
title(slide, "01  /  The problem", "Bidding is easy to demo. Correct bidding is hard to trust.", "A marketplace becomes a distributed-systems problem the moment multiple people bid at once.")
panel(slide, 0.7, 2.15, 5.7, 3.75)
text(slide, 1.05, 2.48, 4.8, 0.35, "What can go wrong?", 21, WHITE, True)
bullet_list(slide, 1.05, 3.05, 4.8, ["Two bidders race on the same price", "A mobile client retries after a timeout", "A stale browser displays an outdated minimum", "A live auction needs updates without refresh"], 16, MUTED)
panel(slide, 6.7, 2.15, 5.9, 3.75, PANEL2)
text(slide, 7.05, 2.48, 5.0, 0.35, "Our design response", 21, WHITE, True)
bullet_list(slide, 7.05, 3.05, 5.0, ["One authoritative bid decision", "Idempotency keys for safe retries", "Realtime event stream for every accepted bid", "State-driven suggestions, never stale hardcoding"], 16, WHITE)
add_footer(slide, 2)

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6]); fill_bg(slide)
title(slide, "02  /  Product", "A live market, not a CRUD form", "The first screen is the actual auction workflow: scan, bid, publish, observe.")
metric(slide, 0.8, 2.1, "LIVE", "SSE event stream", CYAN)
metric(slide, 3.75, 2.1, "1 click", "Publish a listing", MINT)
metric(slide, 6.7, 2.1, "₹", "Minimum bid guard", GOLD)
metric(slide, 9.65, 2.1, "0", "Page refreshes", CYAN)
panel(slide, 0.8, 3.75, 11.5, 2.15)
text(slide, 1.15, 4.05, 3, 0.3, "Judge-visible flow", 19, WHITE, True)
bullet_list(slide, 1.15, 4.55, 10.2, ["Open the same auction in two tabs", "Place a valid bid and watch both views update", "Try a low bid and see the backend reject it", "Publish a new item, bid on it, then remove the listing"], 16, MUTED, 0.39)
add_footer(slide, 3)

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6]); fill_bg(slide)
title(slide, "03  /  Architecture", "A simple path from command to shared truth", "The demo is intentionally explainable: every important behavior has one clear owner.")
# architecture nodes
nodes = [(0.75, 2.55, 2.25, 1.1, "BROWSER\nNext.js UI", CYAN), (3.45, 2.55, 2.45, 1.1, "API\nGo HTTP service", MINT), (6.35, 2.55, 2.5, 1.1, "AUCTION STATE\nSerialized writes", GOLD), (9.25, 2.55, 3.0, 1.1, "EVENT STREAM\nSSE to all clients", CYAN)]
for x,y,w,h,label,accent in nodes:
    panel(slide, x,y,w,h,PANEL2)
    text(slide, x+0.15,y+0.24,w-0.3,0.55,label,15,accent,True,PP_ALIGN.CENTER)
for x in [3.0, 6.0, 8.95]:
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(3.1), Inches(x+0.4), Inches(3.1))
    line.line.color.rgb = MUTED; line.line.width = Pt(2)
text(slide, 1.0, 4.35, 11.2, 0.35, "POST bid  →  validate minimum  →  mutate auction  →  emit accepted event", 21, WHITE, True, PP_ALIGN.CENTER)
panel(slide, 1.1, 5.15, 11.0, 0.85)
text(slide, 1.4, 5.42, 10.4, 0.28, "Production path: PostgreSQL transaction + Redis event distribution + horizontally scaled API instances", 15, MUTED, False, PP_ALIGN.CENTER)
add_footer(slide, 4)

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[6]); fill_bg(slide)
title(slide, "04  /  Correctness", "The bid transaction is the product", "The UI is fast because the rules are strict where they matter: at the write boundary.")
panel(slide, 0.8, 2.05, 5.65, 4.25)
text(slide, 1.15, 2.38, 4.8, 0.3, "Acceptance sequence", 20, WHITE, True)
steps = [("01", "Lock auction state"), ("02", "Check idempotency key"), ("03", "Verify active window"), ("04", "Require current + increment"), ("05", "Write bid and emit event")]
for i,(n,s) in enumerate(steps):
    y=2.95+i*0.58
    text(slide,1.15,y,0.45,0.25,n,11,CYAN,True)
    text(slide,1.8,y,4.1,0.3,s,15,WHITE)
panel(slide, 6.8, 2.05, 5.6, 4.25, PANEL2)
text(slide, 7.15, 2.38, 4.8, 0.3, "Why idempotency matters", 20, WHITE, True)
text(slide, 7.15, 2.95, 4.7, 1.0, "If the client retries the same logical bid, the same key returns the original result. No duplicate bid. No double winner.", 18, MUTED)
panel(slide, 7.15, 4.35, 4.75, 1.15, BG)
text(slide, 7.45, 4.62, 4.2, 0.28, "same key  →  same bid ID", 20, MINT, True, PP_ALIGN.CENTER)
add_footer(slide, 5)

# Slide 6
slide = prs.slides.add_slide(prs.slide_layouts[6]); fill_bg(slide)
title(slide, "05  /  Seller loop", "Anyone can create a market in seconds", "Publishing is part of the live workflow, not a separate admin afterthought.")
panel(slide, 0.8, 2.1, 3.45, 3.6)
text(slide, 1.1, 2.45, 2.8, 0.3, "1  /  Publish", 19, CYAN, True)
text(slide, 1.1, 3.05, 2.7, 1.2, "Title\nCategory\nStarting price\nDuration\nOptional image", 17, WHITE)
panel(slide, 4.95, 2.1, 3.45, 3.6, PANEL2)
text(slide, 5.25, 2.45, 2.8, 0.3, "2  /  Compete", 19, MINT, True)
text(slide, 5.25, 3.05, 2.7, 1.2, "Dynamic minimum\nLive bid suggestions\nCross-tab updates\nRetry-safe bidding", 17, WHITE)
panel(slide, 9.1, 2.1, 3.45, 3.6)
text(slide, 9.4, 2.45, 2.8, 0.3, "3  /  Manage", 19, GOLD, True)
text(slide, 9.4, 3.05, 2.7, 1.2, "Select listing\nDelete listing\nClear feedback\nNo page refresh", 17, WHITE)
add_footer(slide, 6)

# Slide 7
slide = prs.slides.add_slide(prs.slide_layouts[6]); fill_bg(slide)
title(slide, "06  /  Live demo", "A 90-second judge walkthrough", "Keep the story visible: one auction, two tabs, one retry, one new listing.")
steps = [
    ("00:00", "Open two tabs", "Same auction, realtime connected"),
    ("00:15", "Place valid bid", "Both tabs update without refresh"),
    ("00:35", "Try low bid", "Backend rejects stale/low amount"),
    ("00:50", "Publish item", "Seller flow creates a live auction"),
    ("01:05", "Repeat same key", "Idempotency prevents duplicate"),
]
for i,(timecode,head,desc) in enumerate(steps):
    y=2.1+i*0.82
    panel(slide, 0.85, y, 1.15, 0.55, PANEL2)
    text(slide, 0.95,y+0.16,0.95,0.2,timecode,11,CYAN,True,PP_ALIGN.CENTER)
    text(slide, 2.35,y+0.02,3.2,0.25,head,18,WHITE,True)
    text(slide, 5.7,y+0.05,6.5,0.25,desc,15,MUTED)
add_footer(slide, 7)

# Slide 8
slide = prs.slides.add_slide(prs.slide_layouts[6]); fill_bg(slide)
title(slide, "07  /  Scale story", "How we get from a demo to 5,000 virtual bidders", "The correctness model stays the same while the infrastructure becomes distributed.")
panel(slide, 0.8, 2.0, 5.55, 4.1)
text(slide, 1.15, 2.35, 4.7, 0.3, "Today: judge-ready demo", 20, CYAN, True)
bullet_list(slide,1.15,2.95,4.7,["Go service with serialized in-memory state","SSE live updates","Validated publish, bid, delete flows","k6 script for API pressure"],16,WHITE,0.52)
panel(slide, 6.8, 2.0, 5.55, 4.1, PANEL2)
text(slide, 7.15, 2.35, 4.7, 0.3, "Next: production scale", 20, MINT, True)
bullet_list(slide,7.15,2.95,4.7,["PostgreSQL row-level transaction","Redis pub/sub or streams","Load balancer + API replicas","Cloud k6 run at 5,000 VUs"],16,WHITE,0.52)
text(slide, 0.9, 6.45, 11.4, 0.3, "We call 5,000 a virtual-user validation target, not 5,000 real people in a browser.", 13, GOLD, True, PP_ALIGN.CENTER)
add_footer(slide, 8)

# Slide 9
slide = prs.slides.add_slide(prs.slide_layouts[6]); fill_bg(slide)
title(slide, "08  /  Close", "Our differentiation is trust at the moment of truth.", "Synora turns a noisy realtime experience into one consistent auction state.")
panel(slide, 0.8, 2.2, 11.7, 2.0, PANEL2)
text(slide, 1.25, 2.65, 10.8, 0.8, "Fast for bidders.\nSafe for retries. Explainable for judges.", 28, WHITE, True, PP_ALIGN.CENTER)
text(slide, 1.25, 3.75, 10.8, 0.3, "The frontend makes the auction feel alive. The backend makes the result believable.", 16, MUTED, False, PP_ALIGN.CENTER)
text(slide, 0.85, 5.2, 11.5, 0.5, "Thank you", 27, CYAN, True, PP_ALIGN.CENTER)
text(slide, 0.85, 5.85, 11.5, 0.3, "Questions? We can show the live bid path now.", 16, MUTED, False, PP_ALIGN.CENTER)
add_footer(slide, 9)

prs.save(OUT)
print(OUT)
